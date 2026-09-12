from __future__ import annotations

import io
import json
from dataclasses import dataclass

from docx import Document

from job_hunter.config import GroqConfig
from job_hunter.groq_client import chat_json
from job_hunter.keyword_utils import find_skills_line_index
from job_hunter.keyword_utils import fit_score as compute_fit_score
from job_hunter.models import ResumeSuggestion

_SYSTEM_PROMPT = (
    "You tailor resumes for a specific job. You will receive a JSON array of resume "
    'paragraph texts, indexed. Return a JSON object {"paragraphs": [...]} with '
    "EXACTLY the same number of strings, in the same order.\n\n"
    "Two different rules apply depending on the paragraph:\n"
    "1. Skills / technology / tools lists (short lines of comma- or bullet-separated "
    "items, often under a heading like Skills, Expertise, or Technologies): you "
    "SHOULD actively add any of the given 'skills to emphasize' that fit naturally "
    "and are not already present. Adding a listed skill here is expected, not "
    "fabrication.\n"
    "2. Narrative paragraphs (summaries, bullet points about a role): you may "
    "rephrase and re-emphasize wording to align with the target role, but do NOT "
    "invent employers, job titles held, dates, degrees, certifications, or metrics "
    "that are not already present in the input.\n\n"
    "Only return a paragraph unchanged if applying the above rules genuinely "
    "produces no useful change for that specific paragraph — do not default to "
    "leaving everything unchanged."
)


@dataclass(slots=True)
class TailorResult:
    changed: bool
    docx_bytes: bytes | None
    suggestion: ResumeSuggestion
    fit_score: int
    mode: str  # "groq" | "rules"


def _extract_paragraphs(document: Document) -> list[tuple[int, str]]:
    result: list[tuple[int, str]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            result.append((index, text))
    return result


def _apply_paragraph_text(document: Document, index: int, new_text: str) -> None:
    paragraph = document.paragraphs[index]
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _tailor_with_groq(
    paragraphs: list[tuple[int, str]],
    job_title: str,
    job_description: str,
    matched_skills: list[str],
    groq_config: GroqConfig,
) -> list[str] | None:
    user_prompt = (
        f"Target job title: {job_title}\n"
        f"Job description:\n{job_description[:4000]}\n\n"
        f"Skills to emphasize where genuinely applicable: {', '.join(matched_skills)}\n\n"
        "Resume paragraphs (JSON array, index: text):\n"
        + json.dumps([{"index": i, "text": t} for i, t in paragraphs])
    )
    result = chat_json(_SYSTEM_PROMPT, user_prompt, groq_config)
    if not result or "paragraphs" not in result:
        return None
    revised = result["paragraphs"]
    if not isinstance(revised, list) or len(revised) != len(paragraphs):
        return None
    if not all(isinstance(item, str) for item in revised):
        return None
    return revised


def _find_skills_paragraph_index(document: Document) -> int | None:
    return find_skills_line_index([p.text for p in document.paragraphs])


def _rule_based_tailor(
    document: Document,
    matched_skills: list[str],
    profile_skills: list[str],
) -> bool:
    emphasized = matched_skills or profile_skills[:5]
    if not emphasized:
        return False
    skills_index = _find_skills_paragraph_index(document)
    if skills_index is None:
        return False
    existing_text = document.paragraphs[skills_index].text
    additions = [skill for skill in emphasized if skill.lower() not in existing_text.lower()]
    if not additions:
        return False
    new_text = existing_text.rstrip().rstrip(",") + ", " + ", ".join(additions)
    _apply_paragraph_text(document, skills_index, new_text)
    return True


def tailor_docx(
    original_bytes: bytes,
    job_title: str,
    job_description: str,
    matched_skills: list[str],
    profile_skills: list[str],
    groq_config: GroqConfig,
) -> TailorResult:
    """Tailor a .docx resume for a target job while preserving its visual design.

    Text is revised (via Groq, falling back to a deterministic skills-list
    edit) by replacing run text in place; paragraph/document styles, theme,
    section properties, and tables are never touched.
    """
    document = Document(io.BytesIO(original_bytes))
    paragraphs = _extract_paragraphs(document)

    revised_texts: list[str] | None = None
    if paragraphs:
        revised_texts = _tailor_with_groq(
            paragraphs, job_title, job_description, matched_skills, groq_config
        )

    changed = False
    if revised_texts is not None:
        mode = "groq"
        for (para_index, original_text), new_text in zip(paragraphs, revised_texts):
            cleaned = new_text.strip()
            if cleaned and cleaned != original_text.strip():
                _apply_paragraph_text(document, para_index, cleaned)
                changed = True
    else:
        mode = "rules"
        changed = _rule_based_tailor(document, matched_skills, profile_skills)

    emphasized = matched_skills[:5] or profile_skills[:3]
    original_text_blob = "\n".join(text for _, text in paragraphs)
    missing = [
        skill for skill in emphasized if skill.lower() not in original_text_blob.lower()
    ]
    suggestion = ResumeSuggestion(
        mode=mode,
        summary=(
            f"Tailored resume for {job_title}, emphasizing "
            f"{', '.join(emphasized) if emphasized else 'your existing skills'}."
        ),
        keywords_to_emphasize=emphasized,
        missing_resume_keywords=missing,
        bullet_suggestions=[],
    )

    if not changed:
        return TailorResult(
            changed=False,
            docx_bytes=None,
            suggestion=suggestion,
            fit_score=compute_fit_score(original_text_blob, emphasized),
            mode=mode,
        )

    buffer = io.BytesIO()
    document.save(buffer)
    final_text_blob = "\n".join(p.text for p in document.paragraphs)
    return TailorResult(
        changed=True,
        docx_bytes=buffer.getvalue(),
        suggestion=suggestion,
        fit_score=compute_fit_score(final_text_blob, emphasized),
        mode=mode,
    )
