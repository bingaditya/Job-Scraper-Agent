from __future__ import annotations

import io
import unittest
from unittest.mock import patch

from docx import Document

from job_hunter.config import GroqConfig
from job_hunter.docx_tailor import tailor_docx


def _build_fixture_docx() -> bytes:
    document = Document()
    document.add_paragraph("Resume")
    summary = document.add_paragraph()
    run = summary.add_run("Experienced software engineer.")
    run.bold = True
    document.add_paragraph("Skills")
    document.add_paragraph("Python, SQL, Git")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxTailorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.original_bytes = _build_fixture_docx()
        self.groq_config = GroqConfig(api_key="fake-key")

    @patch("job_hunter.docx_tailor.chat_json")
    def test_groq_path_revises_paragraphs_and_preserves_formatting(self, mock_chat_json) -> None:
        mock_chat_json.return_value = {
            "paragraphs": [
                "Resume",
                "Experienced software engineer skilled in Azure and cloud automation.",
                "Skills",
                "Python, SQL, Git",
            ]
        }

        result = tailor_docx(
            original_bytes=self.original_bytes,
            job_title="Cloud Engineer",
            job_description="Looking for Azure experience.",
            matched_skills=["Azure"],
            profile_skills=["Python", "Azure"],
            groq_config=self.groq_config,
        )

        self.assertTrue(result.changed)
        self.assertEqual(result.mode, "groq")
        self.assertIsNotNone(result.docx_bytes)

        rebuilt = Document(io.BytesIO(result.docx_bytes))
        texts = [p.text for p in rebuilt.paragraphs]
        self.assertIn("Azure and cloud automation", texts[1])
        # Bold formatting on the first run of the revised paragraph is preserved.
        self.assertTrue(rebuilt.paragraphs[1].runs[0].bold)

    @patch("job_hunter.docx_tailor.chat_json")
    def test_falls_back_to_rule_based_when_groq_unavailable(self, mock_chat_json) -> None:
        mock_chat_json.return_value = None  # simulates rate limit / failure

        result = tailor_docx(
            original_bytes=self.original_bytes,
            job_title="Backend Engineer",
            job_description="Needs strong Docker skills.",
            matched_skills=["Docker"],
            profile_skills=["Python", "Docker"],
            groq_config=self.groq_config,
        )

        self.assertTrue(result.changed)
        self.assertEqual(result.mode, "rules")
        rebuilt = Document(io.BytesIO(result.docx_bytes))
        skills_paragraph_text = rebuilt.paragraphs[3].text
        self.assertIn("Docker", skills_paragraph_text)
        self.assertIn("Python", skills_paragraph_text)  # original content retained

    @patch("job_hunter.docx_tailor.chat_json")
    def test_no_edit_made_when_nothing_to_add(self, mock_chat_json) -> None:
        mock_chat_json.return_value = None

        result = tailor_docx(
            original_bytes=self.original_bytes,
            job_title="Python Developer",
            job_description="Python role.",
            matched_skills=["Python"],  # already present in the skills paragraph
            profile_skills=["Python"],
            groq_config=self.groq_config,
        )

        self.assertFalse(result.changed)
        self.assertIsNone(result.docx_bytes)


if __name__ == "__main__":
    unittest.main()
