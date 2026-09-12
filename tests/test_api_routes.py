from __future__ import annotations

import io
import unittest
import uuid
from datetime import UTC, datetime
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from job_hunter.api.routes_resume import get_groq_config, get_resume_store, router
from job_hunter.auth import AuthUser, get_current_user
from job_hunter.config import GroqConfig
from job_hunter.resume_store import ResumeNotFoundError
from fastapi import FastAPI

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _build_docx_bytes(skills_text: str = "Python, SQL") -> bytes:
    document = Document()
    document.add_paragraph("Resume")
    document.add_paragraph("Summary paragraph.")
    document.add_paragraph("Skills")
    document.add_paragraph(skills_text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeResumeStore:
    """In-memory stand-in for ResumeStore, matching its public interface."""

    def __init__(self) -> None:
        self.rows: dict[str, dict] = {}
        self.blobs: dict[str, bytes] = {}
        self.profiles: dict[str, dict] = {}

    def get_candidate_profile(self, user_id: str) -> dict | None:
        return self.profiles.get(user_id)

    def upsert_candidate_profile(self, user_id: str, profile, source_version_id: str) -> dict:
        row = {
            "user_id": user_id,
            "skills": profile.skills,
            "preferred_titles": profile.preferred_titles,
            "experience_years": profile.experience_years,
            "summary": profile.summary,
            "source_version_id": source_version_id,
            "extraction_mode": profile.mode,
        }
        self.profiles[user_id] = row
        return row

    def _next_version_number(self, user_id: str) -> int:
        existing = [r for r in self.rows.values() if r["user_id"] == user_id]
        return max((r["version_number"] for r in existing), default=0) + 1

    def list_versions(self, user_id: str) -> list[dict]:
        rows = [r for r in self.rows.values() if r["user_id"] == user_id]
        return sorted(rows, key=lambda r: r["created_at"], reverse=True)

    def get_version(self, user_id: str, version_id: str) -> dict:
        row = self.rows.get(version_id)
        if row is None or row["user_id"] != user_id:
            raise ResumeNotFoundError(version_id)
        return row

    def get_current_version(self, user_id: str) -> dict | None:
        for row in self.rows.values():
            if row["user_id"] == user_id and row["is_current"]:
                return row
        return None

    def download_bytes(self, user_id: str, version_id: str) -> tuple[bytes, dict]:
        row = self.get_version(user_id, version_id)
        return self.blobs[version_id], row

    def _store(self, user_id: str, file_bytes: bytes, **fields) -> dict:
        for row in self.rows.values():
            if row["user_id"] == user_id:
                row["is_current"] = False
        version_id = str(uuid.uuid4())
        row = {
            "id": version_id,
            "user_id": user_id,
            "version_number": self._next_version_number(user_id),
            "is_current": True,
            "storage_path": f"{user_id}/{version_id}.docx",
            "file_size_bytes": len(file_bytes),
            "mime_type": DOCX_MIME,
            "created_at": datetime.now(tz=UTC).isoformat(),
            "job_title": None,
            "tailoring_mode": None,
            "fit_score": None,
            **fields,
        }
        self.rows[version_id] = row
        self.blobs[version_id] = file_bytes
        return row

    def create_upload_version(self, user_id: str, file_bytes: bytes, original_filename: str) -> dict:
        return self._store(
            user_id, file_bytes, source="upload", original_filename=original_filename
        )

    def create_tailored_version(
        self,
        user_id: str,
        source_version: dict,
        file_bytes: bytes,
        job_title: str,
        job_description: str,
        job_id: str | None,
        tailoring_mode: str,
        suggestion,
        fit_score: int,
    ) -> dict:
        return self._store(
            user_id,
            file_bytes,
            source="ai_tailored",
            original_filename=source_version["original_filename"],
            job_title=job_title,
            tailoring_mode=tailoring_mode,
            fit_score=fit_score,
        )


def _build_test_app(store: FakeResumeStore, user_id: str = "user-1") -> FastAPI:
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_current_user] = lambda: AuthUser(id=user_id, email="t@example.com")
    app.dependency_overrides[get_resume_store] = lambda: store
    app.dependency_overrides[get_groq_config] = lambda: GroqConfig(api_key=None)
    return app


class ApiRoutesTests(unittest.TestCase):
    def test_upload_list_and_download_round_trip(self) -> None:
        store = FakeResumeStore()
        client = TestClient(_build_test_app(store))
        docx_bytes = _build_docx_bytes()

        upload_resp = client.post(
            "/api/resumes",
            files={"file": ("resume.docx", docx_bytes, DOCX_MIME)},
        )
        self.assertEqual(upload_resp.status_code, 201)
        version_id = upload_resp.json()["id"]
        self.assertEqual(upload_resp.json()["version_number"], 1)
        self.assertTrue(upload_resp.json()["is_current"])

        list_resp = client.get("/api/resumes")
        self.assertEqual(list_resp.status_code, 200)
        self.assertEqual(len(list_resp.json()), 1)

        download_resp = client.get(f"/api/resumes/{version_id}/download")
        self.assertEqual(download_resp.status_code, 200)
        self.assertEqual(download_resp.content, docx_bytes)  # byte-identical

    @patch("job_hunter.resume_profile.chat_json")
    def test_upload_extracts_candidate_profile(self, mock_chat_json) -> None:
        mock_chat_json.return_value = {
            "skills": ["Python", "Docker"],
            "preferred_titles": ["Backend Engineer"],
            "experience_years": 5,
            "summary": "Backend engineer with Python and Docker experience.",
        }
        store = FakeResumeStore()
        client = TestClient(_build_test_app(store))
        client.post(
            "/api/resumes",
            files={"file": ("resume.docx", _build_docx_bytes(), DOCX_MIME)},
        )
        profile = store.get_candidate_profile("user-1")
        self.assertIsNotNone(profile)
        self.assertEqual(profile["skills"], ["Python", "Docker"])
        self.assertEqual(profile["extraction_mode"], "groq")

    def test_rejects_non_docx_upload(self) -> None:
        store = FakeResumeStore()
        client = TestClient(_build_test_app(store))
        resp = client.post(
            "/api/resumes",
            files={"file": ("notes.txt", b"hello", "text/plain")},
        )
        self.assertEqual(resp.status_code, 400)

    @patch("job_hunter.docx_tailor.chat_json")
    def test_tailor_creates_new_version_and_keeps_previous(self, mock_chat_json) -> None:
        mock_chat_json.return_value = {
            "paragraphs": ["Resume", "Tailored summary paragraph.", "Skills", "Python, SQL"]
        }
        store = FakeResumeStore()
        client = TestClient(_build_test_app(store))
        docx_bytes = _build_docx_bytes()
        upload_resp = client.post(
            "/api/resumes",
            files={"file": ("resume.docx", docx_bytes, DOCX_MIME)},
        )
        original_version_id = upload_resp.json()["id"]

        tailor_resp = client.post(
            "/api/resumes/tailor",
            json={"job_title": "Data Engineer", "job_description": "Needs Python.", "skills": ["Python"]},
        )
        self.assertEqual(tailor_resp.status_code, 200)
        body = tailor_resp.json()
        self.assertTrue(body["changed"])
        self.assertEqual(body["version"]["version_number"], 2)
        self.assertTrue(body["version"]["is_current"])

        # Previous version retained and no longer current.
        list_resp = client.get("/api/resumes")
        versions = {v["id"]: v for v in list_resp.json()}
        self.assertEqual(len(versions), 2)
        self.assertFalse(versions[original_version_id]["is_current"])

    def test_tailor_without_upload_returns_404(self) -> None:
        store = FakeResumeStore()
        client = TestClient(_build_test_app(store))
        resp = client.post(
            "/api/resumes/tailor",
            json={"job_title": "Data Engineer", "job_description": "Needs Python."},
        )
        self.assertEqual(resp.status_code, 404)


if __name__ == "__main__":
    unittest.main()
